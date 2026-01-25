"""
Excel Report Generator - Converts JSON compliance report to Excel workbook
Generates multi-sheet Excel with gap analysis, guidelines, and policies
"""

import json
import logging
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class ComplianceExcelGenerator:
    """Generate Excel reports from JSON compliance data"""
    
    def __init__(self, gap_analysis_path, guidelines_path, policy_path, excel_output_path):
        """
        Initialize Excel generator
        
        Args:
            gap_analysis_path: Path to gap_analysis.json
            guidelines_path: Path to current_guidelines.json
            policy_path: Path to current_policy.json
            excel_output_path: Path where Excel should be saved
        """
        self.gap_analysis_path = gap_analysis_path
        self.guidelines_path = guidelines_path
        self.policy_path = policy_path
        self.excel_path = excel_output_path
        self.wb = Workbook()
        self.wb.remove(self.wb.active)  # Remove default sheet
        
        # Define styles
        self.header_fill = PatternFill(start_color="1a3a52", end_color="1a3a52", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        self.subheader_fill = PatternFill(start_color="2d5a7b", end_color="2d5a7b", fill_type="solid")
        self.subheader_font = Font(bold=True, color="FFFFFF", size=10)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self.center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        self.left_align = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        self.gap_analysis = None
        self.guidelines = None
        self.policy = None
    
    def load_data(self):
        """Load all required JSON files"""
        try:
            with open(self.gap_analysis_path, 'r', encoding='utf-8') as f:
                self.gap_analysis = json.load(f)
            logger.info(f"Loaded gap analysis from: {self.gap_analysis_path}")
            
            with open(self.guidelines_path, 'r', encoding='utf-8') as f:
                self.guidelines = json.load(f)
            logger.info(f"Loaded guidelines from: {self.guidelines_path}")
            
            with open(self.policy_path, 'r', encoding='utf-8') as f:
                self.policy = json.load(f)
            logger.info(f"Loaded policy from: {self.policy_path}")
            
            return True
        except FileNotFoundError as e:
            logger.error(f"File not found: {e}")
            return False
        except json.JSONDecodeError as e:
            logger.error(f"Error parsing JSON: {e}")
            return False
    
    def _add_header_row(self, ws, headers):
        """Add header row to worksheet"""
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.center_align
            cell.border = self.border
    
    def _auto_adjust_columns(self, ws):
        """Auto-adjust column widths"""
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def create_summary_sheet(self):
        """Create summary/overview sheet"""
        ws = self.wb.create_sheet("Summary")
        
        # Title
        ws['A1'] = "RBI COMPLIANCE MONITORING REPORT"
        ws['A1'].font = Font(bold=True, size=14, color="1a3a52")
        ws.merge_cells('A1:D1')
        ws['A1'].alignment = self.left_align
        
        # Summary metrics
        row = 3
        summary_data = [
            ("Generated Date", datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ("Total Guidelines Analyzed", str(self.gap_analysis.get('total_guidelines_analyzed', 0))),
            ("Total Policy Clauses", str(self.gap_analysis.get('total_policy_clauses', 0))),
            ("", ""),
        ]
        
        status_summary = self.gap_analysis.get('status_summary', {})
        summary_data.extend([
            ("COMPLIANCE STATUS SUMMARY", ""),
            ("Covered", str(status_summary.get('covered', 0))),
            ("Outdated", str(status_summary.get('outdated', 0))),
            ("Missing", str(status_summary.get('missing', 0))),
        ])
        
        for label, value in summary_data:
            cell_label = ws.cell(row=row, column=1)
            cell_label.value = label
            cell_label.font = Font(bold=True) if label else Font()
            cell_label.border = self.border
            
            cell_value = ws.cell(row=row, column=2)
            cell_value.value = value
            cell_value.border = self.border
            row += 1
        
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
    
    def create_gap_analysis_sheet(self):
        """Create detailed gap analysis sheet"""
        ws = self.wb.create_sheet("Gap Analysis")
        
        headers = ["Guideline ID", "Guideline Title", "Status", "Keyword Count", "RBI Link"]
        self._add_header_row(ws, headers)
        
        gaps = self.gap_analysis.get('gaps', [])
        for row_num, gap in enumerate(gaps, 2):
            ws.cell(row=row_num, column=1).value = gap.get('guideline_id', '')
            ws.cell(row=row_num, column=2).value = gap.get('guideline_title', '')
            
            status = gap.get('compliance_status', '')
            status_cell = ws.cell(row=row_num, column=3)
            status_cell.value = status
            
            # Color code status
            if status == "COVERED":
                status_cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            elif status == "OUTDATED":
                status_cell.fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
            elif status == "MISSING":
                status_cell.fill = PatternFill(start_color="FFB6C6", end_color="FFB6C6", fill_type="solid")
            
            ws.cell(row=row_num, column=4).value = gap.get('keyword_count', '')
            ws.cell(row=row_num, column=5).value = gap.get('guideline_link', '')
            
            # Apply borders and alignment
            for col in range(1, 6):
                cell = ws.cell(row=row_num, column=col)
                cell.border = self.border
                cell.alignment = self.left_align
        
        self._auto_adjust_columns(ws)
    
    def create_guidelines_sheet(self):
        """Create RBI guidelines sheet"""
        ws = self.wb.create_sheet("RBI Guidelines")
        
        headers = ["Guideline ID", "Title", "Link", "Date", "Description"]
        self._add_header_row(ws, headers)
        
        entries = self.guidelines.get('entries', [])
        for row_num, entry in enumerate(entries, 2):
            ws.cell(row=row_num, column=1).value = row_num - 1
            ws.cell(row=row_num, column=2).value = entry.get('title', '')
            ws.cell(row=row_num, column=3).value = entry.get('link', '')
            ws.cell(row=row_num, column=4).value = entry.get('published_date', '')
            
            desc = entry.get('description', '')
            if desc and len(desc) > 200:
                desc = desc[:200] + "..."
            ws.cell(row=row_num, column=5).value = desc
            
            # Apply borders and alignment
            for col in range(1, 6):
                cell = ws.cell(row=row_num, column=col)
                cell.border = self.border
                cell.alignment = self.left_align
            
            ws.row_dimensions[row_num].height = 30
        
        self._auto_adjust_columns(ws)
    
    def create_policy_sheet(self):
        """Create bank policy sheet"""
        ws = self.wb.create_sheet("Bank Policies")
        
        headers = ["Clause ID", "Clause Text"]
        self._add_header_row(ws, headers)
        
        clauses = self.policy.get('clauses', [])
        for row_num, clause in enumerate(clauses, 2):
            ws.cell(row=row_num, column=1).value = clause.get('clause_id', '')
            ws.cell(row=row_num, column=2).value = clause.get('clause_text', '')
            
            # Apply borders and alignment
            for col in range(1, 3):
                cell = ws.cell(row=row_num, column=col)
                cell.border = self.border
                cell.alignment = self.left_align
            
            ws.row_dimensions[row_num].height = 25
        
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 80
    
    def create_compliance_mapping_sheet(self):
        """Create guideline-to-policy mapping sheet"""
        ws = self.wb.create_sheet("Compliance Mapping")
        
        headers = ["Guideline", "Guideline Title", "Status", "Related Policies", "Match Score"]
        self._add_header_row(ws, headers)
        
        gaps = self.gap_analysis.get('gaps', [])
        for row_num, gap in enumerate(gaps, 2):
            ws.cell(row=row_num, column=1).value = gap.get('guideline_id', '')
            ws.cell(row=row_num, column=2).value = gap.get('guideline_title', '')
            
            status = gap.get('compliance_status', '')
            status_cell = ws.cell(row=row_num, column=3)
            status_cell.value = status
            
            # Color code
            if status == "COVERED":
                status_cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            elif status == "MISSING":
                status_cell.fill = PatternFill(start_color="FFB6C6", end_color="FFB6C6", fill_type="solid")
            
            ws.cell(row=row_num, column=4).value = "See policies sheet"
            ws.cell(row=row_num, column=5).value = gap.get('match_percentage', 'N/A')
            
            # Apply borders
            for col in range(1, 6):
                cell = ws.cell(row=row_num, column=col)
                cell.border = self.border
                cell.alignment = self.center_align
        
        self._auto_adjust_columns(ws)
    
    def generate_excel(self):
        """Generate the complete Excel workbook"""
        if not self.load_data():
            return False
        
        try:
            logger.info("Creating Summary sheet...")
            self.create_summary_sheet()
            
            logger.info("Creating Gap Analysis sheet...")
            self.create_gap_analysis_sheet()
            
            logger.info("Creating RBI Guidelines sheet...")
            self.create_guidelines_sheet()
            
            logger.info("Creating Bank Policies sheet...")
            self.create_policy_sheet()
            
            logger.info("Creating Compliance Mapping sheet...")
            self.create_compliance_mapping_sheet()
            
            # Save workbook
            self.wb.save(str(self.excel_path))
            logger.info(f"Excel workbook saved to: {self.excel_path}")
            return True
        
        except Exception as e:
            logger.error(f"Error generating Excel: {e}", exc_info=True)
            return False


def main():
    """Main function"""
    import sys
    
    # Setup paths
    project_root = Path(__file__).parent.parent.parent
    gap_analysis = project_root / "data" / "output" / "gap_analysis.json"
    guidelines = project_root / "data" / "rbi_guidelines" / "current_guidelines.json"
    policy = project_root / "data" / "bank_policies" / "current_policy.json"
    excel_report = project_root / "reports" / "compliance_report.xlsx"
    
    # Generate Excel
    logger.info("Starting Excel report generation...")
    generator = ComplianceExcelGenerator(gap_analysis, guidelines, policy, excel_report)
    
    if generator.generate_excel():
        logger.info(f"[SUCCESS] Excel report generated successfully")
        logger.info(f"Output file: {excel_report}")
        print(f"\n{'='*80}")
        print(f"EXCEL REPORT GENERATED SUCCESSFULLY")
        print(f"{'='*80}")
        print(f"Report Location: {excel_report}")
        print(f"File Size: {excel_report.stat().st_size / 1024:.2f} KB")
        print(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"Sheets: Summary, Gap Analysis, RBI Guidelines, Bank Policies, Compliance Mapping")
        print(f"{'='*80}\n")
        return 0
    else:
        logger.error("[FAILED] Excel generation failed")
        print(f"\n{'='*80}")
        print(f"EXCEL REPORT GENERATION FAILED")
        print(f"{'='*80}\n")
        return 1


if __name__ == "__main__":
    import sys
    sys.exit(main())
