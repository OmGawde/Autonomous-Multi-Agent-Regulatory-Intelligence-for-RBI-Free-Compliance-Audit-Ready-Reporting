import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('Precisa-BSA-Sample-Report.xlsx', data_only=True)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n{'='*80}")
    print(f"SHEET: {sheet_name}")
    print(f"Dimensions: {ws.dimensions} | Max Row: {ws.max_row} | Max Col: {ws.max_column}")
    print(f"{'='*80}")
    
    max_rows = min(ws.max_row, 60)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), 1):
        row_data = [str(cell) if cell is not None else "" for cell in row]
        if any(cell.strip() for cell in row_data):
            # Truncate long rows
            display = ' | '.join(row_data[:20])
            if len(display) > 500:
                display = display[:500] + "..."
            print(f"R{i}: {display}")
