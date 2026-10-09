"""
The downloadable Excel analysis report.

The workbook is a deliverable someone reads and acts on, so the tests check the
two things that would make it misleading rather than merely ugly: that its
headline figures reproduce what the statement itself prints, and that every
formula uses a function Excel can actually evaluate.
"""

import csv
import json
import sys
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
OUTPUTS = REPO_ROOT / "Feature Extraction" / "outputs"
BACKEND = REPO_ROOT / "Website" / "backend"

for p in (str(REPO_ROOT), str(BACKEND)):
    if p not in sys.path:
        sys.path.insert(0, p)

pytest.importorskip("openpyxl", reason="openpyxl is required to build the report")

# What the Axis statement prints about itself, independent of our parsing.
AXIS_PRINTED = {"opening": 29155.31, "closing": 17707.93,
                "debits": 774966.38, "credits": 763519.00, "rows": 422}

# Every function used must be Excel-2007-era: LibreOffice and older Excel both
# evaluate these without a prefix, and openpyxl writes formulas verbatim.
SAFE_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "IFERROR", "IF"}


@pytest.fixture(scope="module")
def workbook():
    features_path = OUTPUTS / "Axis_features.json"
    classified_path = OUTPUTS / "Axis_classified.csv"
    if not features_path.exists() or not classified_path.exists():
        pytest.skip("Axis pipeline outputs not present; run the pipeline first")

    from services.report_builder import build_report

    features = json.loads(features_path.read_text(encoding="utf-8"))
    with open(classified_path, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    if len(rows) != AXIS_PRINTED["rows"]:
        pytest.skip(f"Axis_classified.csv holds {len(rows)} rows, expected {AXIS_PRINTED['rows']}")

    return build_report(features, rows, {"applicant_name": "MILIND SURESH WALANJU",
                                         "bank_name": "Axis"})


@pytest.fixture(scope="module")
def icici_workbook():
    """ICICI is the statement that actually has investments and loans detected."""
    features_path = OUTPUTS / "ICICI_features.json"
    classified_path = OUTPUTS / "ICICI_classified.csv"
    if not features_path.exists() or not classified_path.exists():
        pytest.skip("ICICI pipeline outputs not present; run the pipeline first")

    from services.report_builder import build_report

    features = json.loads(features_path.read_text(encoding="utf-8"))
    with open(classified_path, encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    return build_report(features, rows, {"applicant_name": "GIRISH BABAJI GIDAYE",
                                         "bank_name": "ICICI"})


def _summary_rows(ws):
    header = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Description")
    months = [c for c in range(2, ws.max_column + 1)
              if ws.cell(header, c).value not in (None, "Total")]
    labels = {ws.cell(r, 1).value: r for r in range(header + 1, ws.max_row + 1)}
    return header, months, labels


def test_every_expected_sheet_is_present(workbook):
    names = set(workbook.sheetnames)
    for expected in ("Contents", "Summary", "Transactions", "Categories", "Counterparty",
                     "Cash Flow", "Daily Balance", "Credit Score", "Verification"):
        assert expected in names, f"missing sheet {expected}"
    engines = [n for n in names if n.startswith("E ")]
    assert len(engines) == 9, f"expected one sheet per engine, got {engines}"


def test_summary_opening_and_closing_match_the_printed_statement(workbook):
    """The figures a reader will check first must agree with the document."""
    ws = workbook["Summary"]
    _header, months, labels = _summary_rows(ws)

    opening = ws.cell(labels["Opening Balance"], months[0]).value
    closing = ws.cell(labels["Closing Balance"], months[-1]).value
    assert opening == pytest.approx(AXIS_PRINTED["opening"], abs=0.01)
    assert closing == pytest.approx(AXIS_PRINTED["closing"], abs=0.01)


def test_monthly_openings_chain_from_the_previous_closing(workbook):
    ws = workbook["Summary"]
    _header, months, labels = _summary_rows(ws)
    for prev, cur in zip(months, months[1:]):
        closing = ws.cell(labels["Closing Balance"], prev).value
        opening = ws.cell(labels["Opening Balance"], cur).value
        assert opening == pytest.approx(closing, abs=0.01), (
            "a month must open where the previous one closed"
        )


def test_summary_totals_sum_to_the_printed_totals(workbook):
    """
    The Total cells are formulas, so compute them here the way Excel will and
    check the result against what the bank printed.
    """
    ws = workbook["Summary"]
    _header, months, labels = _summary_rows(ws)

    for label, expected in (("Debit Txns", AXIS_PRINTED["debits"]),
                            ("Credit Txns", AXIS_PRINTED["credits"])):
        row = labels[label]
        total = sum(ws.cell(row, c).value or 0 for c in months)
        assert total == pytest.approx(expected, abs=0.01), label

    counted = sum(ws.cell(labels["Debit Txns Count"], c).value or 0 for c in months) + \
        sum(ws.cell(labels["Credit Txns Count"], c).value or 0 for c in months)
    assert counted == AXIS_PRINTED["rows"]


def test_totals_are_formulas_not_baked_constants(workbook):
    """A report that cannot recalculate is a screenshot with extra steps."""
    ws = workbook["Summary"]
    header, months, labels = _summary_rows(ws)
    total_col = months[-1] + 1
    value = ws.cell(labels["Debit Txns"], total_col).value
    assert isinstance(value, str) and value.startswith("=SUM(")


def test_every_formula_uses_a_function_excel_can_evaluate(workbook):
    """
    Post-2007 names need an `_xlfn.` prefix and the spilling ones cannot be
    written by openpyxl at all; either way the cell renders as #NAME?.
    """
    import re

    pattern = re.compile(r"([A-Z_][A-Z0-9_.]*)\s*\(")
    offenders = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    for fn in pattern.findall(cell.value.upper()):
                        if fn not in SAFE_FUNCTIONS:
                            offenders.append(f"{ws.title}!{cell.coordinate}: {fn}")
    assert not offenders, f"unsupported functions: {offenders[:10]}"


def test_no_formula_points_at_an_empty_range(workbook):
    import re

    from openpyxl.utils import range_boundaries

    pattern = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
    empty = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                for m in pattern.finditer(cell.value):
                    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
                    min_c, _, max_c, _ = range_boundaries(f"{c1}{r1}:{c2}{r2}")
                    if r2 < r1 or max_c < min_c:
                        empty.append(f"{ws.title}!{cell.coordinate}")
    assert not empty, f"formulas over inverted/empty ranges: {empty[:10]}"


def test_no_formula_operand_is_a_text_cell(workbook):
    """
    The bug this exists to prevent: a "-" placeholder written into a cell that a
    formula then reads. IFERROR does not catch it, because text is not an error
    -- it flows straight through and `1000 - "-"` renders as #VALUE!. The credit
    score sheet shipped that way once.
    """
    import re

    from openpyxl.utils import column_index_from_string, range_boundaries

    range_ref = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
    single_ref = re.compile(r"(?<![A-Z0-9:$!])\$?([A-Z]{1,3})\$?(\d+)(?![0-9(:])")

    offenders = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                refs = []
                for m in range_ref.finditer(cell.value):
                    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
                    min_c, _, max_c, _ = range_boundaries(f"{c1}{r1}:{c2}{r2}")
                    refs += [(rr, cc) for rr in range(r1, r2 + 1)
                             for cc in range(min_c, max_c + 1)]
                for m in single_ref.finditer(range_ref.sub("", cell.value)):
                    refs.append((int(m.group(2)), column_index_from_string(m.group(1))))

                for rr, cc in refs:
                    target = ws.cell(rr, cc).value
                    if isinstance(target, str) and not target.startswith("="):
                        offenders.append(
                            f"{ws.title}!{cell.coordinate} reads "
                            f"{ws.cell(rr, cc).coordinate}={target!r}"
                        )

    assert not offenders, f"formulas reading text cells: {offenders[:8]}"


@pytest.mark.parametrize("bank", ["Axis", "ICICI", "SBI", "Bank_Of_India", "BOB", "Union"])
def test_score_itemisation_reconciles_to_the_reported_score(bank):
    """
    The breakdown re-derives each band from src/feature_aggregator.py. If the
    two ever disagree the sheet would show a decomposition that does not add up
    to the score printed above it.
    """
    from services.report_builder import _score_steps

    path = OUTPUTS / f"{bank}_features.json"
    if not path.exists():
        pytest.skip(f"{bank} pipeline output not present")

    features = json.loads(path.read_text(encoding="utf-8"))
    reported = (features.get("underwriting_summary") or {}).get("credit_score")
    if reported is None:
        pytest.skip(f"{bank} has no credit score")

    itemised = 1000 + sum(step[3] for step in _score_steps(features))
    assert itemised == reported, (
        f"{bank}: steps sum to {itemised} but the engine reported {reported}"
    )


def test_score_adjustments_are_numbers_not_placeholders(workbook):
    """Every adjustment cell must be arithmetic-safe."""
    ws = workbook["Credit Score"]
    header = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Step")
    final = next(r for r in range(header, ws.max_row + 1) if ws.cell(r, 1).value == "Final score")
    checked = 0
    for r in range(header + 1, final):
        value = ws.cell(r, 4).value
        assert isinstance(value, (int, float)), (
            f"row {r} adjustment is {value!r}; a text cell here yields #VALUE!"
        )
        checked += 1
    assert checked >= 6, "expected the starting score plus one row per scoring band"


def test_no_cell_contains_a_raw_python_repr(workbook):
    """
    Nested values must be rendered as real tables, not stringified. The
    investment engine's `investments_detected` is a dict of dicts and arrived
    in a single cell as `{'count': 10, 'total_debit': 27207.84, ...}` --
    unreadable, and impossible to sum, sort or filter.
    """
    offenders = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                text = cell.value
                if isinstance(text, str) and (
                    text.startswith("{'") or text.startswith("[{") or text.startswith("['")
                ):
                    offenders.append(f"{ws.title}!{cell.coordinate}: {text[:50]}")
    assert not offenders, f"raw Python reprs in cells: {offenders[:6]}"


def test_nested_dicts_become_tables_with_one_column_per_field(icici_workbook):
    """`investments_detected` must have Count / Total debit / Total credit columns."""
    ws = icici_workbook["E · Investment"]
    header = next(
        (r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Item"), None
    )
    if header is None:
        pytest.skip("this statement detected no investments")

    headers = [ws.cell(header, c).value for c in range(1, 6)]
    assert "Count" in headers and "Total debit" in headers

    # And the numeric columns get a live total.
    totals = [r for r in range(header, ws.max_row + 1) if ws.cell(r, 1).value == "Total"]
    assert totals, "a table of numbers needs a total row"
    value = ws.cell(totals[0], 2).value
    assert isinstance(value, str) and value.startswith("=SUM(")


def test_number_formats_suit_what_the_figure_is(icici_workbook):
    """A score is not currency, and a fractional frequency is not an integer."""
    ws = icici_workbook["E · Investment"]
    formats = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if isinstance(label, str):
            formats[label] = ws.cell(r, 2).number_format

    if "Investment ratio" in formats:
        assert formats["Investment ratio"] == "0.0%"
    if "Long term investment score" in formats:
        assert formats["Long term investment score"] == "0.00"
    if "Sip count" in formats:
        assert "0" in formats["Sip count"] and "." not in formats["Sip count"]


def test_verification_sheet_carries_the_fidelity_verdict(workbook):
    """The differentiator over the reference format -- do not let it go missing."""
    ws = workbook["Verification"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "VERIFIED" in text
    assert "debits_match_printed_total" in text


def test_report_does_not_reference_the_source_vendor(workbook):
    """
    The layout was modelled on a competitor's sample report. The name must not
    appear anywhere in what we ship.
    """
    text = " ".join(
        str(c.value) for ws in workbook.worksheets
        for row in ws.iter_rows() for c in row if c.value
    ).lower()
    assert "precisa" not in text
    assert "precisa" not in " ".join(workbook.sheetnames).lower()
